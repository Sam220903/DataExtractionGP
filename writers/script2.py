import json
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Configuración inicial ──────────────────────────────────────────────────
VOTE_CODE = {"Favor": 1, "Contra": 2, "Abstencion": 3, "NA": 0, "Secreta": 4}

VOTE_COLORS = {
    1: "C6EFCE",   # Verde claro  – A favor
    2: "FFC7CE",   # Rojo claro   – En contra
    3: "F8CBAD",   # Naranja      – Abstención
    4: "BFBFBF",   # Gris medio   – Secreta
    0: "FFFFFF",   # Blanco       – Ausencia / Sin voto
}

# Paletas del formato original
SALMON      = "E6B8AF"
RED_HEADER  = "FF0000"
WHITE       = "FFFFFF"
GREY_EMPTY  = "CCCCCC"

# Estilos reutilizables de celda
def _fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

thin = Side(style="thin", color="DDDDDD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_TOP    = Alignment(vertical="top")
ALIGN_WRAP   = Alignment(wrap_text=True, vertical="bottom")


def parse_date(date_str):
    try:
        return datetime.strptime(date_str, "%d/%m/%y")
    except Exception:
        return datetime.min


# ── Carga y procesamiento de datos ──────────────────────────────────────────
with open("../data/votaciones.json", "r", encoding="utf-8") as f:
    data = json.load(f)

registros = data["registros"]
registros.sort(key=lambda x: (parse_date(x.get("fecha", "")), x.get("numero_votacion", 0)))

# Lista ordenada de diputados con su partido
deputies_info = {}   # nombre -> partido
for r in registros:
    for v in r.get("votaciones", []):
        name = v["diputado"].strip()
        if name not in deputies_info:
            deputies_info[name] = v.get("partido", "")

deputies_list = list(deputies_info.keys())

# ── Creación del libro ───────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Hoja 1"

output_filename = "Votaciones.xlsx"
output_folder = "../outputs"

NUM_VOTES = len(registros)   # columnas de votaciones
NUM_DEPS  = len(deputies_list)

# Columnas: A=etiqueta, B=partido, C...(C+NUM_VOTES-1)=votaciones
VOTE_START_COL = 3   # columna C

# ── FILA 1: folio ────────────────────────────────────────────────────────────
ws.cell(row=1, column=1, value="folio")

# ── FILA 2: num  (fondo rojo en A) ──────────────────────────────────────────
c = ws.cell(row=2, column=1, value="num")
c.fill = _fill(RED_HEADER)

for col_idx, r in enumerate(registros, start=VOTE_START_COL):
    cell = ws.cell(row=2, column=col_idx, value=r.get("numero_votacion", ""))
    cell.fill = _fill(WHITE)
    cell.alignment = ALIGN_CENTER
    cell.font = Font(bold=True)

# ── FILA 3: fecha ────────────────────────────────────────────────────────────
ws.cell(row=3, column=1, value="fecha")

for col_idx, r in enumerate(registros, start=VOTE_START_COL):
    fecha_str = r.get("fecha", "")
    fecha_val = parse_date(fecha_str) if fecha_str else None
    cell = ws.cell(row=3, column=col_idx, value=fecha_val if fecha_val and fecha_val != datetime.min else fecha_str)
    cell.number_format = "DD/MM/YY"
    cell.alignment = ALIGN_CENTER

# ── FILA 4: Sesión  (fondo salmón) ──────────────────────────────────────────
c = ws.cell(row=4, column=1, value="Sesión")
c.fill = _fill(SALMON)

for col_idx, r in enumerate(registros, start=VOTE_START_COL):
    cell = ws.cell(row=4, column=col_idx, value=r.get("sesion", ""))
    cell.fill = _fill(SALMON)
    cell.alignment = ALIGN_CENTER

# ── FILA 5: Votación  (fondo salmón) ────────────────────────────────────────
c = ws.cell(row=5, column=1, value="Votación")
c.fill = _fill(SALMON)

for col_idx, r in enumerate(registros, start=VOTE_START_COL):
    cell = ws.cell(row=5, column=col_idx, value=r.get("numero_votacion", ""))
    cell.fill = _fill(SALMON)
    cell.alignment = ALIGN_CENTER

# ── FILA 6: tema  (altura ajustada) ─────────────────────────────────────────
ws.cell(row=6, column=1, value="tema")
ws.row_dimensions[6].height = 20.25

for col_idx, r in enumerate(registros, start=VOTE_START_COL):
    cell = ws.cell(row=6, column=col_idx, value=r.get("tema", ""))
    cell.fill = _fill(WHITE)
    cell.alignment = ALIGN_WRAP
    cell.font = Font(color="313131")

# ── FILA 7: IN (1) PA (2)  (fondo salmón) ───────────────────────────────────
c = ws.cell(row=7, column=1, value="IN (1) PA (2) ")
c.fill = _fill(SALMON)

for col_idx, r in enumerate(registros, start=VOTE_START_COL):
    sesion = r.get("sesion", "").strip().lower()
    tipo_val = 1 if sesion == "previa" else 2
    cell = ws.cell(row=7, column=col_idx, value=tipo_val)
    cell.fill = _fill(SALMON)

# ── FILA 8: Periodo  (fondo salmón) ─────────────────────────────────────────
c = ws.cell(row=8, column=1, value="Periodo")
c.fill = _fill(SALMON)

for col_idx, r in enumerate(registros, start=VOTE_START_COL):
    cell = ws.cell(row=8, column=col_idx, value=r.get("periodo", 1))
    cell.fill = _fill(SALMON)

# ── FILA 9: encabezado Diputado / Partido ────────────────────────────────────
c_dip = ws.cell(row=9, column=1, value="Diputado")
c_dip.fill = _fill(RED_HEADER)
c_dip.font = Font(bold=True)
c_dip.alignment = ALIGN_TOP

c_par = ws.cell(row=9, column=2, value="Partido")
c_par.fill = _fill(SALMON)
c_par.alignment = Alignment(vertical="bottom")

# ── FILAS 10+: Diputados y sus votos ─────────────────────────────────────────
DEP_START_ROW = 10

for row_offset, deputy in enumerate(deputies_list):
    current_row = DEP_START_ROW + row_offset
    partido = deputies_info[deputy]

    # Col A: nombre
    dep_cell = ws.cell(row=current_row, column=1, value=deputy)
    dep_cell.alignment = Alignment(vertical="top")

    # Col B: partido
    par_cell = ws.cell(row=current_row, column=2, value=partido if partido else None)

    # Cols de votos
    for col_idx, r in enumerate(registros, start=VOTE_START_COL):
        tipo_vot = r.get("tipo_votacion", "")
        votos_map = {v["diputado"].strip(): v["voto"].capitalize() for v in r.get("votaciones", [])}

        if tipo_vot == "Secreta":
            val = VOTE_CODE["Secreta"]
        elif deputy in votos_map:
            val = VOTE_CODE.get(votos_map[deputy], VOTE_CODE["NA"])
        else:
            val = None

        v_cell = ws.cell(row=current_row, column=col_idx, value=val)
        v_cell.alignment = ALIGN_CENTER

        if val is None:
            v_cell.value = "-"
            v_cell.fill = _fill(GREY_EMPTY)   # Gris = sin dato / licencia
        else:
            color = VOTE_COLORS.get(val, WHITE)
            v_cell.fill = _fill(color)

# ── FILAS FOOTER: totales ────────────────────────────────────────────────────
FOOTER_START = DEP_START_ROW + NUM_DEPS + 1   # una fila de separación

footer_labels = [
    ("Votos a favor",    "a_favor"),
    ("Votos en contra",  "en_contra"),
    ("Votos en secreto", "en_secreto"),
    ("Abstención",       "abstenciones"),
    ("Ausencia",         "ausencias"),
    ("Sumatoria total",  "total_votos"),
]

for i, (label, key) in enumerate(footer_labels):
    row = FOOTER_START + i
    label_cell = ws.cell(row=row, column=2, value=label)
    if label == "Sumatoria total":
        label_cell.font = Font(bold=True)
    if label == "Votos en secreto":
        label_cell.fill = _fill(GREY_EMPTY)

    for col_idx, r in enumerate(registros, start=VOTE_START_COL):
        val = r.get(key, 0)
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.alignment = ALIGN_CENTER

# Leyenda debajo
legend_row = FOOTER_START + len(footer_labels) + 2
ws.cell(row=legend_row,     column=1, value="NO SE PUEDEN CONTABILIZAR PORQUE NO CORRESPONDEL AL ACTA O NO SE ENCONTRÓ EL EXCEL").fill = _fill("980000")
ws.cell(row=legend_row + 1, column=1, value="VOTACIÓN SECRETA").fill = _fill("B7B7B7")
ws.cell(row=legend_row + 2, column=1, value="NO SE PUEDEN CONTABILIZAR PORQUE NO EXISTE EN EL ACTA").fill  = _fill("660000")

# ── Anchos de columna ────────────────────────────────────────────────────────
ws.column_dimensions["A"].width = 36.38
ws.column_dimensions["B"].width = 13.0
for col_idx in range(VOTE_START_COL, VOTE_START_COL + NUM_VOTES):
    ws.column_dimensions[get_column_letter(col_idx)].width = 13.0

# ── Inmovilizar paneles: columnas A+B fijas, filas 1-9 fijas ─────────────────
# El original congela en "C10" (= col C hacia la derecha, fila 10 hacia abajo)
ws.freeze_panes = "C10"

# ── Guardar ──────────────────────────────────────────────────────────────────
output_path = os.path.join(output_folder, output_filename)
wb.save(output_path)
print(f"Archivo generado: {output_path}")