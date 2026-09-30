#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera "Sesiones.xlsx" a partir de uno o más JSON, con el mismo formato que el
Excel original ("Copia de Sesiones.xlsx").

Estructura de carpetas (obligatoria):
    proyecto/
    ├── writer/    <- este script
    ├── data/      <- JSON de entrada
    └── outputs/   <- Excel generado

Uso:
    python generar_excel.py
    python generar_excel.py --data-dir RUTA --output-dir RUTA --output-name Otro.xlsx
    python generar_excel.py --verbose        (muestra todos los avisos, sin recortar)

Las columnas derivadas (K, L, M, P, S, T, U) son fórmulas de Excel: se recalculan
solas si alguien edita las columnas de entrada.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import Counter, OrderedDict, defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# =============================================================================
# CONFIGURACIÓN  (lo único que normalmente hay que tocar)
# =============================================================================

# Rutas: se resuelven desde la ubicación de ESTE archivo, no desde el directorio
# de trabajo, para que funcione igual lo lance quien lo lance.
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_DATA_DIR = SCRIPT_DIR.parent / "data"
DEFAULT_OUTPUT_DIR = SCRIPT_DIR.parent / "outputs"

# Nombre del Excel de salida: sin fecha ni versión; cada ejecución lo sobrescribe.
# (Original "Copia de Sesiones.xlsx" -> se quita "Copia de" -> "Sesiones.xlsx")
OUTPUT_NAME = "Sesiones.xlsx"

# Un JSON por hoja. "json" es el archivo dentro de data/, "sheet" el nombre de la hoja.
SHEETS = [
    {"json": "sesiones.json", "sheet": "Hoja 1"},
]

# Opciones
DAYFIRST = True                    # 10/09/2024 = 10 de septiembre
NULL_TEXTS = {"", "null", "none", "nan", "n/a", "na", "nat"}   # textos que equivalen a vacío
MAX_RETRASO_MIN_AVISO = 180        # retrasos mayores a esto se reportan como atípicos
WRAP_OBSERVACIONES = False         # el original no ajusta el texto de Observaciones
CORREGIR_ENCABEZADO_BLANCO = False # el original tiene el encabezado M en blanco sobre amarillo
EXTRA_COLUMNS_STYLE = {"head_fill": "E5B8B7", "fill": "FFFFFF", "h": "center"}  # columnas nuevas del JSON
FREEZE_PANES = "A2"
DEFAULT_COL_WIDTH = 12.63
DEFAULT_ROW_HEIGHT = 15.75
MAX_EJEMPLOS_AVISO = 8             # ejemplos por tipo de aviso (con --verbose se muestran todos)

# Nota y vínculo de la celda A1 (el original tiene un comentario en hilo + hipervínculo)
A1_COMMENT = ("La información se obtiene de las Gacetas Legislativas mensuales "
              "(https://www.congresopuebla.gob.mx/index.php?option=com_content&view=article&id=12868) "
              "a partir de: órden del día y actas de sesiones.")
A1_COMMENT_AUTHOR = "Autor"
A1_HYPERLINK = "http://www.reportelegislativo.com.mx/prueba/s2"

# Colores del original
YELLOW, PINK, GREY, LGREY, RED, GREEN, WHITE = (
    "FFFF00", "E5B8B7", "B7B7B7", "CCCCCC", "FF0000", "00FF00", "FFFFFF")

# -----------------------------------------------------------------------------
# Definición de columnas (orden = orden en el Excel).
#   id       nombre corto para las fórmulas ({id} se reemplaza por la celda de esa fila)
#   key      clave del JSON (se compara normalizada: sin acentos/mayúsculas/espacios/signos)
#   header   texto del encabezado en Excel
#   kind     text | number | date | time | formula
#   formula  plantilla de la fórmula (solo kind="formula")
#   fmt      formato de número     head_fill/head_font/fill/h/wrap = estilo
# -----------------------------------------------------------------------------
COLUMNS = [
    dict(id="num", key="Número de sesión", header="Número de sesión", kind="number", fmt="General",
         head_fill=YELLOW, head_font="0000FF", fill=WHITE, h="center", wrap=True, width=8.38),
    dict(id="letra", key="Sesión letra", header="Sesión letra", kind="text", fmt="General",
         head_fill=YELLOW, fill=WHITE, h="center", wrap=True),
    dict(id="anio", key="Año legislatura", header="Año legislatura", kind="text", fmt="General",
         head_fill=PINK, fill=GREY, h="center", wrap=True),
    dict(id="periodo", key="Periodo legislatura", header="Periodo Legislatura", kind="text", fmt="General",
         head_fill=PINK, fill=GREY, h="center", wrap=True, width=13.88),
    dict(id="nper", key="Número del periodo", header="Número del periodo", kind="number", fmt="General",
         head_fill=PINK, fill=GREY, h="center", wrap=True),
    dict(id="tipo", key="Tipo de sesión", header="Tipo de sesión", kind="text", fmt="@",
         head_fill=PINK, fill=GREY, h="center", wrap=True),
    dict(id="obs", key="Observaciones", header="Observaciones:", kind="text", fmt="General",
         head_fill=WHITE, fill=None, h=None, wrap=None, font=("Arial", 10), optional=True),
    dict(id="fecha", key="Fecha", header="Fecha", kind="date", fmt="d/MM/yyyy",
         head_fill=PINK, fill=GREY, h="center", wrap=None),
    dict(id="gaceta", key="Fecha Gaceta Parlamentaria", header="Fecha Gaceta Parlamentaria", kind="text", fmt="@",
         head_fill=YELLOW, fill=WHITE, h="center", wrap=True),
    dict(id="cita", key="Hora de cita para inicio de sesión", header="Hora de cita para inicio de sesión",
         kind="time", fmt="h:mm:ss AM/PM", head_fill=PINK, fill=LGREY, h="center", wrap=None),
    dict(id="hret", key="Hora retraso de inicio de sesión", header="Hora retraso de inicio de sesión",
         kind="formula", fmt="hh:mm:ss", head_fill=YELLOW, fill=YELLOW, h="center", wrap=None,
         formula='IF(AND(ISNUMBER({cita}),ISNUMBER({ini})),{ini}-{cita},"")'),
    dict(id="retmin", key="Retraso de inicio de sesión (minutos)", header="Retraso de inicio de sesión (minutos)",
         kind="formula", fmt="General", head_fill=YELLOW, fill=None, h="center", wrap=None,
         formula='IF(ISNUMBER({hret}),ROUND({hret}*1440,0),"")'),
    dict(id="ret", key="Retraso de inicio de sesión", header="Retraso de inicio de sesión",
         kind="formula", fmt="h:mm:ss", head_fill=YELLOW, head_font="FFFFFF", fill=RED, h="center", wrap=None,
         formula='IF(ISNUMBER({hret}),{hret},"")'),
    dict(id="ini", key="Inicio de sesión", header="Inicio de sesión", kind="time", fmt="h:mm:ss",
         head_fill=PINK, fill=GREY, h="right", wrap=None),
    dict(id="fin", key="Fin de sesión", header="Fin de sesión", kind="time", fmt="h:mm:ss",
         head_fill=PINK, fill=GREY, h="right", wrap=None),
    dict(id="dur", key="Duración", header="Duración", kind="formula", fmt="h:mm:ss",
         head_fill=YELLOW, fill=None, h="center", wrap=None,
         formula='IF(AND(ISNUMBER({ini}),ISNUMBER({fin})),MOD({fin}-{ini},1),"")'),
    dict(id="prog", key="Total de asuntos programados por sesión", header="Total de asuntos programados por sesión",
         kind="number", fmt="General", head_fill=YELLOW, fill=GREY, h="center", wrap=None),
    dict(id="abo", key="Total de asuntos abordados por sesión", header="Total de asuntos abordados por sesión",
         kind="number", fmt="General", head_fill=PINK, fill=GREEN, h="center", wrap=None),
    dict(id="noabo", key="Total de asuntos no abordados por sesión", header="Total de asuntos no abordados por sesión",
         kind="formula", fmt="0", head_fill=YELLOW, fill=WHITE, h="center", wrap=True,
         formula='IF(AND(ISNUMBER({prog}),ISNUMBER({abo})),{prog}-{abo},"")'),
    dict(id="pct", key="Porcentaje de asuntos abordados por sesión", header="Porcentaje de asuntos abordados por sesión",
         kind="formula", fmt="0%", head_fill=YELLOW, fill=WHITE, h="center", wrap=None,
         formula='IF(AND(ISNUMBER({prog}),ISNUMBER({abo})),IF({prog}=0,"",{abo}/{prog}),"")'),
    dict(id="coef", key="Coeficiente de asuntos abordados por sesión", header="Coeficiente de asuntos abordados por sesión",
         kind="formula", fmt="0.00", head_fill=YELLOW, fill=WHITE, h="center", wrap=True,
         formula='IF(AND(ISNUMBER({prog}),ISNUMBER({abo})),IF({prog}=0,"",{abo}/{prog}),"")'),
]

# =============================================================================
# UTILIDADES
# =============================================================================

def norm_key(s) -> str:
    """Normaliza una clave: sin acentos, minúsculas, sin espacios ni signos."""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "", s.lower())


def is_null(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip().lower() in NULL_TEXTS) \
        or (isinstance(v, float) and v != v)


MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
         "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}

_DATE_FORMATS_DAY = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%Y-%m-%d", "%Y/%m/%d"]
_DATE_FORMATS_MONTH = ["%m/%d/%Y", "%m-%d-%Y", "%m/%d/%y", "%Y-%m-%d", "%Y/%m/%d"]
_TIME_RE = re.compile(
    r"^(\d{1,3}):(\d{2})(?::(\d{2})(?:[.,]\d+)?)?\s*(?:([ap])\s*\.?\s*m\.?)?$", re.I)


def parse_date(v):
    """-> date | None (None = no interpretable)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 <= v <= 80000:                      # serial de Excel
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
        return None
    if not isinstance(v, str):
        return None
    s = v.strip()
    s = re.split(r"[T ]\d{1,2}:\d{2}", s)[0].strip()    # descarta parte horaria
    m = re.match(r"^(\d{1,2})\s+(?:de\s+)?([a-záéíóú]+)\s+(?:de\s+)?(\d{4})$", s, re.I)
    if m and m.group(2).lower() in MESES:
        try:
            return date(int(m.group(3)), MESES[m.group(2).lower()], int(m.group(1)))
        except ValueError:
            return None
    for f in (_DATE_FORMATS_DAY if DAYFIRST else _DATE_FORMATS_MONTH):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def clock_seconds(v):
    """Segundos desde 00:00 para textos tipo '10:00:00 a.m.', '22:15', '0:35:00'.
    Admite horas >= 24 (duraciones). -> int | None."""
    if isinstance(v, time):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(round(v * 86400)) if 0 <= v < 1 else None
    if not isinstance(v, str):
        return None
    m = _TIME_RE.match(v.strip())
    if not m:
        return None
    h, mi, se, ap = int(m.group(1)), int(m.group(2)), int(m.group(3) or 0), m.group(4)
    if mi > 59 or se > 59:
        return None
    if ap:
        if not 1 <= h <= 12:
            return None
        h = (h % 12) + (12 if ap.lower() == "p" else 0)
    return h * 3600 + mi * 60 + se


def parse_time(v):
    """-> datetime.time | None (una hora del día; >= 24 h no es válida)."""
    sec = clock_seconds(v)
    if sec is None or sec >= 86400:
        return None
    return time(sec // 3600, (sec % 3600) // 60, sec % 60)


def parse_number(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v == int(v) else v
    if isinstance(v, str):
        s = v.strip().replace(",", ".") if re.fullmatch(r"-?\d+,\d+", v.strip()) else v.strip()
        try:
            f = float(s)
        except ValueError:
            return None
        return int(f) if f == int(f) else f
    return None


def convert(kind: str, raw):
    """-> (valor, estado) con estado 'ok' | 'empty' | 'invalid'.
    Si es 'invalid' el valor devuelto es el original (se conserva tal cual)."""
    if is_null(raw):
        return None, "empty"
    if isinstance(raw, (dict, list)):
        return json.dumps(raw, ensure_ascii=False), "ok"
    if kind == "date":
        d = parse_date(raw)
        return (datetime(d.year, d.month, d.day), "ok") if d else (raw, "invalid")
    if kind == "time":
        t = parse_time(raw)
        return (t, "ok") if t is not None else (raw, "invalid")
    if kind == "number":
        n = parse_number(raw)
        return (n, "ok") if n is not None else (raw, "invalid")
    return (raw.strip() if isinstance(raw, str) else raw), "ok"      # text


def guess_kind(values) -> str:
    """Para columnas nuevas del JSON: decide date / time / number / text según los valores."""
    vals = [v for v in values if not is_null(v)]
    if not vals:
        return "text"
    for kind, fn in (("date", parse_date), ("time", parse_time), ("number", parse_number)):
        if kind == "time" and not all(isinstance(v, str) and ":" in v for v in vals):
            continue
        if kind == "date" and not all(isinstance(v, str) for v in vals):
            continue
        if all(fn(v) is not None for v in vals):
            return kind
    return "text"


# =============================================================================
# AVISOS
# =============================================================================

class Avisos:
    def __init__(self):
        self.grupos = OrderedDict()

    def add(self, categoria: str, detalle: str):
        self.grupos.setdefault(categoria, []).append(detalle)

    def total(self):
        return sum(len(v) for v in self.grupos.values())

    def imprimir(self, verbose=False):
        if not self.grupos:
            print("\nAvisos: ninguno. Los datos cuadran.")
            return
        print(f"\nAVISOS ({self.total()} en {len(self.grupos)} categorías)")
        print("=" * 72)
        for cat, items in self.grupos.items():
            print(f"\n• {cat}  [{len(items)}]")
            mostrar = items if verbose else items[:MAX_EJEMPLOS_AVISO]
            for it in mostrar:
                print(f"    - {it}")
            if len(mostrar) < len(items):
                print(f"    ... y {len(items) - len(mostrar)} más (usa --verbose para verlos todos)")


# =============================================================================
# LECTURA DE JSON
# =============================================================================

def load_records(path: Path, avisos: Avisos):
    with open(path, "r", encoding="utf-8-sig") as fh:
        data = json.load(fh)
    if isinstance(data, dict):
        listas = [v for v in data.values() if isinstance(v, list) and v and all(isinstance(x, dict) for x in v)]
        if len(listas) == 1:
            data = listas[0]
        elif data and all(isinstance(v, dict) for v in data.values()):
            data = list(data.values())
        else:
            raise ValueError(f"{path.name}: el JSON es un objeto y no encuentro una lista de registros.")
    if not isinstance(data, list):
        raise ValueError(f"{path.name}: se esperaba una lista de registros.")
    records = []
    for i, r in enumerate(data, start=1):
        if not isinstance(r, dict):
            avisos.add("Registros que no son objetos (omitidos)", f"{path.name}: posición {i}: {type(r).__name__}")
            continue
        records.append(r)
    return records


def build_layout(records, avisos: Avisos, json_name: str):
    """Empareja las claves del JSON con las columnas definidas.
    Devuelve (columnas, mapa_clave_original_por_id)."""
    # Claves reales del JSON (orden de aparición) agrupadas por forma normalizada
    reales = OrderedDict()
    for r in records:
        for k in r:
            reales.setdefault(norm_key(k), [])
            if k not in reales[norm_key(k)]:
                reales[norm_key(k)].append(k)
    for nk, ks in reales.items():
        if len(ks) > 1:
            avisos.add("Claves duplicadas tras normalizar (se usa la primera)", f"{json_name}: {ks}")

    columnas, mapa, usadas = [], {}, set()
    for spec in COLUMNS:
        c = dict(spec)
        nk = norm_key(spec["key"])
        if nk in reales:
            mapa[c["id"]] = reales[nk][0]
            usadas.add(nk)
        else:
            mapa[c["id"]] = None
            avisos.add("Columnas que no vienen en el JSON (se dejan vacías)",
                       f"{json_name}: '{spec['key']}'"
                       + (" — es fórmula, se calcula igual" if spec["kind"] == "formula" else ""))
        columnas.append(c)

    # Claves nuevas -> columnas al final
    for nk, ks in reales.items():
        if nk in usadas:
            continue
        clave = ks[0]
        kind = guess_kind([r.get(clave) for r in records])
        c = dict(id=f"extra_{nk}", key=clave, header=str(clave).strip(), kind=kind, extra=True,
                 fmt={"date": "d/MM/yyyy", "time": "h:mm:ss"}.get(kind, "General"),
                 head_fill=EXTRA_COLUMNS_STYLE["head_fill"], fill=EXTRA_COLUMNS_STYLE["fill"],
                 h=EXTRA_COLUMNS_STYLE["h"], wrap=True)
        columnas.append(c)
        mapa[c["id"]] = clave
        avisos.add("Columnas nuevas en el JSON (agregadas al final)", f"{json_name}: '{clave}' (tipo detectado: {kind})")
    return columnas, mapa


# =============================================================================
# CONSTRUCCIÓN DE LA HOJA
# =============================================================================

def argb(color: str) -> str:
    """'FF0000' -> 'FFFF0000' (canal alfa opaco, como en el original)."""
    return "FF" + color if len(color) == 6 else color


THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(cell, col):
    font_color = col.get("head_font", "000000")
    if col["id"] == "ret" and CORREGIR_ENCABEZADO_BLANCO:
        font_color = "000000"
    cell.font = Font(name="Calibri", size=11, bold=True, color=argb(font_color))
    cell.fill = PatternFill("solid", fgColor=argb(col["head_fill"]))
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)
    cell.number_format = "@"


def style_body(cell, col):
    name, size = col.get("font", ("Calibri", 11))
    cell.font = Font(name=name, size=size, color=argb("000000"))   # negro explícito (igual que el original)
    if col.get("fill"):
        cell.fill = PatternFill("solid", fgColor=argb(col["fill"]))
    cell.border = BORDER
    wrap = col.get("wrap")
    if col["id"] == "obs" and WRAP_OBSERVACIONES:
        wrap = True
    if col.get("h") or wrap:
        cell.alignment = Alignment(horizontal=col.get("h"), wrap_text=wrap or None)
    cell.number_format = col.get("fmt", "General")


def fila_label(row, rec_vals, columnas_por_id):
    letra = rec_vals.get("letra")
    fecha = rec_vals.get("fecha")
    f = fecha.strftime("%d/%m/%Y") if isinstance(fecha, datetime) else (fecha or "sin fecha")
    return f"fila {row} ({letra or 'sin nombre de sesión'}, {f})"


def secs(t):
    return t.hour * 3600 + t.minute * 60 + t.second if isinstance(t, time) else None


def build_sheet(wb, first: bool, sheet_name: str, records, json_name: str, avisos: Avisos):
    ws = wb.active if first else wb.create_sheet()
    ws.title = sheet_name
    columnas, mapa = build_layout(records, avisos, json_name)
    letra_de = {c["id"]: get_column_letter(i) for i, c in enumerate(columnas, start=1)}

    ws.sheet_format.defaultColWidth = DEFAULT_COL_WIDTH
    ws.sheet_format.defaultRowHeight = DEFAULT_ROW_HEIGHT
    ws.sheet_format.customHeight = True

    # Encabezados
    for i, c in enumerate(columnas, start=1):
        cell = ws.cell(row=1, column=i, value=c["header"])
        style_header(cell, c)
        if c.get("width"):
            ws.column_dimensions[get_column_letter(i)].width = c["width"]

    # A1: comentario + hipervínculo (como el original)
    if A1_COMMENT:
        com = Comment(A1_COMMENT, A1_COMMENT_AUTHOR)
        com.width, com.height = 320, 120
        ws["A1"].comment = com
    if A1_HYPERLINK:
        ws["A1"].hyperlink = A1_HYPERLINK
        ws["A1"].font = Font(name="Calibri", size=11, bold=True, color=argb("0000FF"))

    # Columnas de entrada que no existen en el JSON (para no repetir avisos derivados)
    ausentes = {cid for cid, key in mapa.items() if key is None}

    # Datos
    vistos_letra = {}
    fechas_previas = None
    faltantes = defaultdict(list)          # id -> [filas con vacío]
    fila = 1
    for idx, rec in enumerate(records, start=1):
        if all(is_null(v) for v in rec.values()):
            avisos.add("Registros vacíos (omitidos)", f"{json_name}: registro {idx}")
            continue
        fila += 1
        vals, raws = {}, {}
        for c in columnas:
            key = mapa.get(c["id"])
            raw = rec.get(key) if key is not None else None
            raws[c["id"]] = raw
            kind = c["kind"]
            if kind == "formula":
                continue
            v, estado = convert(kind, raw)
            vals[c["id"]] = v
            if estado == "invalid":
                avisos.add("Valores no interpretables (se conservan tal cual)",
                           f"{json_name} fila {fila}: columna '{c['header']}' = {raw!r}")
        etiqueta = fila_label(fila, vals, None)

        refs = {cid: f"{letter}{fila}" for cid, letter in letra_de.items()}
        for i, c in enumerate(columnas, start=1):
            cell = ws.cell(row=fila, column=i)
            if c["kind"] == "formula":
                cell.value = "=" + c["formula"].format(**refs)
            else:
                cell.value = vals[c["id"]]
                if vals[c["id"]] is None and mapa.get(c["id"]) is not None:
                    faltantes[c["id"]].append(fila)
            style_body(cell, c)

        # --- Chequeos de coherencia (avisos) -------------------------------
        cita, ini, fin = vals.get("cita"), vals.get("ini"), vals.get("fin")
        if isinstance(ini, time) and isinstance(fin, time) and secs(fin) < secs(ini):
            avisos.add("Fin de sesión anterior al inicio (¿cruza medianoche o hay un error?)",
                       f"{etiqueta}: inicio {ini}, fin {fin}")
        if isinstance(cita, time) and isinstance(ini, time):
            if secs(ini) < secs(cita):
                avisos.add("Inicio de sesión anterior a la hora de cita (retraso negativo)",
                           f"{etiqueta}: cita {cita}, inicio {ini}")
            elif (secs(ini) - secs(cita)) / 60 > MAX_RETRASO_MIN_AVISO:
                avisos.add(f"Retrasos atípicos (> {MAX_RETRASO_MIN_AVISO} min)",
                           f"{etiqueta}: cita {cita}, inicio {ini} = {(secs(ini) - secs(cita)) // 60} min")
        prog, abo = vals.get("prog"), vals.get("abo")
        if isinstance(prog, (int, float)) and isinstance(abo, (int, float)):
            if abo > prog:
                avisos.add("Asuntos abordados > programados", f"{etiqueta}: programados {prog}, abordados {abo}")
            if prog == 0:
                avisos.add("Asuntos programados = 0 (porcentaje y coeficiente quedan vacíos)", etiqueta)
        letra = vals.get("letra")
        if letra:
            if letra in vistos_letra:
                avisos.add("Nombre de sesión repetido",
                           f"'{letra}' en fila {vistos_letra[letra]} y en fila {fila}")
            else:
                vistos_letra[letra] = fila
        f = vals.get("fecha")
        if isinstance(f, datetime):
            if fechas_previas and f < fechas_previas[0]:
                avisos.add("Fechas fuera de orden cronológico (menor que la fila anterior)",
                           f"{etiqueta} es anterior a la fila {fechas_previas[1]} "
                           f"({fechas_previas[0].strftime('%d/%m/%Y')})")
            fechas_previas = (f, fila)

        # --- Valores derivados del JSON vs. lo que calcularán las fórmulas ---
        _chequear_derivados(vals, raws, etiqueta, avisos, ausentes)

    ultima = max(fila, 1)

    # Resumen de datos faltantes por columna
    total = ultima - 1
    for c in columnas:
        filas = faltantes.get(c["id"], [])
        if not filas or c.get("optional"):
            continue
        if len(filas) == total:
            avisos.add("Columnas vacías en todos los registros", f"'{c['header']}' no tiene ningún dato en el JSON")
        else:
            avisos.add("Datos faltantes por columna",
                       f"'{c['header']}': {len(filas)} vacíos (filas {_resumen_filas(filas)})")

    # Filtro, panel congelado
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{ultima}"
    ws.freeze_panes = FREEZE_PANES
    return ws, total, len(columnas)


def _resumen_filas(filas, n=10):
    txt = ", ".join(str(x) for x in filas[:n])
    return txt + (", …" if len(filas) > n else "")


def fmt_hms(sec):
    signo = "-" if sec < 0 else ""
    sec = abs(int(sec))
    return f"{signo}{sec // 3600}:{(sec % 3600) // 60:02d}:{sec % 60:02d}"


# Columnas de entrada de las que depende cada columna derivada
DEPENDENCIAS = {"hret": ("cita", "ini"), "ret": ("cita", "ini"), "retmin": ("cita", "ini"),
                "dur": ("ini", "fin"), "noabo": ("prog", "abo"), "pct": ("prog", "abo"),
                "coef": ("prog", "abo")}


def _chequear_derivados(vals, raws, etiqueta, avisos: Avisos, ausentes=frozenset()):
    """Compara los valores derivados que trae el JSON con lo que darán las fórmulas."""
    cita, ini, fin = vals.get("cita"), vals.get("ini"), vals.get("fin")
    prog, abo = vals.get("prog"), vals.get("abo")
    esperados = {}
    if isinstance(cita, time) and isinstance(ini, time):
        esperados["hret"] = secs(ini) - secs(cita)
        esperados["ret"] = esperados["hret"]
        esperados["retmin"] = round(esperados["hret"] / 60)
    if isinstance(ini, time) and isinstance(fin, time):
        esperados["dur"] = (secs(fin) - secs(ini)) % 86400
    if isinstance(prog, (int, float)) and isinstance(abo, (int, float)):
        esperados["noabo"] = prog - abo
        if prog != 0:
            esperados["pct"] = abo / prog * 100
            esperados["coef"] = abo / prog
    nombres = {"hret": "Hora retraso", "ret": "Retraso", "retmin": "Retraso (minutos)", "dur": "Duración",
               "noabo": "No abordados", "pct": "Porcentaje", "coef": "Coeficiente"}
    for cid, nombre in nombres.items():
        raw = raws.get(cid)
        if cid in ("hret", "ret", "dur"):
            json_v = None if is_null(raw) else clock_seconds(raw)
            tol = 1
        else:
            json_v = None if is_null(raw) else parse_number(raw)
            tol = {"pct": 0.06, "coef": 0.0006}.get(cid, 0.5)
        exp = esperados.get(cid)
        if json_v is None and exp is None:
            continue
        if any(dep in ausentes for dep in DEPENDENCIAS[cid]):
            continue          # ya se avisó que falta esa columna de entrada
        if json_v is None:
            avisos.add("Derivados: el JSON no lo trae pero la fórmula sí lo calcula", f"{etiqueta}: {nombre}")
        elif exp is None:
            avisos.add("Derivados: el JSON trae valor pero faltan datos de entrada para la fórmula",
                       f"{etiqueta}: {nombre} = {raw!r}")
        elif abs(json_v - exp) > tol:
            avisos.add("Derivados: el valor del JSON no cuadra con la fórmula",
                       f"{etiqueta}: {nombre}: JSON={raw!r}, la fórmula dará "
                       + (fmt_hms(exp) if cid in ("hret", "ret", "dur") else f"{round(exp, 4)}"))


# =============================================================================
# PRINCIPAL
# =============================================================================

def parse_args():
    p = argparse.ArgumentParser(description="Genera el Excel de sesiones a partir de los JSON.")
    p.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR, help=f"carpeta de JSON (por defecto {DEFAULT_DATA_DIR})")
    p.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR, help=f"carpeta de salida (por defecto {DEFAULT_OUTPUT_DIR})")
    p.add_argument("--output-name", default=OUTPUT_NAME, help=f"nombre del Excel (por defecto {OUTPUT_NAME})")
    p.add_argument("--verbose", action="store_true", help="mostrar todos los avisos sin recortar")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    data_dir = args.data_dir.expanduser().resolve()
    out_dir = args.output_dir.expanduser().resolve()
    out_path = out_dir / args.output_name
    avisos = Avisos()

    if not data_dir.is_dir():
        print(f"ERROR: no existe la carpeta de datos: {data_dir}")
        return 1
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        print(f"ERROR: no pude crear la carpeta de salida {out_dir}: {e}")
        return 1

    wb = Workbook()
    resumen = []
    for n, hoja in enumerate(SHEETS):
        ruta = data_dir / hoja["json"]
        if not ruta.is_file():
            existentes = ", ".join(sorted(p.name for p in data_dir.glob("*.json"))) or "(ninguno)"
            print(f"ERROR: no encuentro {ruta}\n       JSON disponibles en {data_dir}: {existentes}")
            return 1
        try:
            records = load_records(ruta, avisos)
        except json.JSONDecodeError as e:
            print(f"ERROR: {ruta.name} no es un JSON válido (línea {e.lineno}, columna {e.colno}): {e.msg}")
            return 1
        except ValueError as e:
            print(f"ERROR: {e}")
            return 1
        if not records:
            print(f"ERROR: {ruta.name} no contiene registros.")
            return 1
        _, filas, ncols = build_sheet(wb, n == 0, hoja["sheet"], records, hoja["json"], avisos)
        resumen.append((hoja["sheet"], hoja["json"], filas, ncols))

    wb.calculation.fullCalcOnLoad = True      # Excel recalcula todas las fórmulas al abrir

    lock = out_path.with_name("~$" + out_path.name)
    if lock.exists():
        print(f"AVISO: existe {lock.name}; el archivo parece estar abierto en Excel.")
    try:
        wb.save(out_path)
    except PermissionError:
        print(f"\nNO SE PUDO GUARDAR: '{out_path.name}' está abierto en Excel (u otro programa).\n"
              f"Ciérralo y vuelve a ejecutar el script.\nRuta: {out_path}")
        return 1
    except OSError as e:
        print(f"ERROR al guardar {out_path}: {e}")
        return 1

    print(f"Excel generado: {out_path}")
    for hoja, js, filas, ncols in resumen:
        print(f"  hoja '{hoja}' <- {js}: {filas} filas de datos, {ncols} columnas")
    avisos.imprimir(verbose=args.verbose)
    return 0


if __name__ == "__main__":
    sys.exit(main())