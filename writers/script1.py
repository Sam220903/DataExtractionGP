import json
import os
import sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Config de estilo (replicando Perfiles_Plantilla.xlsx)
# ---------------------------------------------------------------------------
FUENTE_ENCABEZADO = Font(name="Calibri", size=11, bold=True)
RELLENO_ENCABEZADO = PatternFill(start_color="FF9CC2E5", end_color="FF9CC2E5", fill_type="solid")
ALINEACION_ENCABEZADO = Alignment(vertical="bottom")

FUENTE_DATO = Font(name="Arial", size=11, color="FF000000")

BORDE_FINO = Border(
    left=Side(style="thin", color="FF000000"),
    right=Side(style="thin", color="FF000000"),
    top=Side(style="thin", color="FF000000"),
    bottom=Side(style="thin", color="FF000000"),
)

VALOR_VACIO = "N/A"  # criterio de la plantilla original para celdas sin dato

ANCHO_MINIMO = 8
ANCHO_MAXIMO = 60  # evita columnas gigantes cuando hay textos muy largos (direcciones, preparación, etc.)


def autoajustar_anchos(ws, ultima_col: int, ultima_fila: int) -> None:
    """
    Recorre cada columna y ajusta su ancho al contenido más largo que tenga,
    dentro de un mínimo y un máximo razonables.
    """
    for col_idx in range(1, ultima_col + 1):
        letra = get_column_letter(col_idx)
        max_len = len(str(ws.cell(row=1, column=col_idx).value or ""))

        for fila_idx in range(2, ultima_fila + 1):
            valor = ws.cell(row=fila_idx, column=col_idx).value
            if valor is not None:
                largo = len(str(valor))
                if largo > max_len:
                    max_len = largo

        ancho = max(ANCHO_MINIMO, min(max_len + 2, ANCHO_MAXIMO))
        ws.column_dimensions[letra].width = ancho


def cargar_registros(ruta_json: str) -> list[dict]:
    """Lee el JSON y devuelve la lista de registros bajo la clave 'registros'."""
    with open(ruta_json, encoding="utf-8") as f:
        data = json.load(f)

    if isinstance(data, dict) and "registros" in data:
        registros = data["registros"]
    elif isinstance(data, list):
        registros = data
    else:
        raise ValueError(
            "Formato de JSON no reconocido: se esperaba un objeto con la clave "
            "'registros' o directamente una lista de registros."
        )

    if not registros:
        raise ValueError("El JSON no contiene registros.")

    return registros


def construir_columnas(registros: list[dict]) -> list[str]:
    """
    Determina el orden de columnas a partir de las claves del primer registro,
    y agrega al final cualquier clave nueva que aparezca en otros registros
    (por si el JSON se actualiza con campos adicionales en el futuro).
    """
    columnas = list(registros[0].keys())
    vistas = set(columnas)

    for registro in registros[1:]:
        for clave in registro.keys():
            if clave not in vistas:
                columnas.append(clave)
                vistas.add(clave)

    return columnas


def normalizar_valor(valor):
    """Aplica el mismo criterio que la plantilla: None -> 'N/A'."""
    if valor is None:
        return VALOR_VACIO
    if isinstance(valor, str) and valor.strip() == "":
        return VALOR_VACIO
    return valor


def generar_excel(registros: list[dict], ruta_salida: str) -> None:
    columnas = construir_columnas(registros)

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja 1"

    # --- Encabezados ---
    # Columna A sin encabezado (igual que la plantilla: es el consecutivo)
    ws.cell(row=1, column=1, value=None)
    for idx, nombre_col in enumerate(columnas, start=2):
        celda = ws.cell(row=1, column=idx, value=nombre_col)
        celda.font = FUENTE_ENCABEZADO
        celda.fill = RELLENO_ENCABEZADO
        celda.alignment = ALINEACION_ENCABEZADO
        celda.border = BORDE_FINO

    # --- Datos ---
    for fila_idx, registro in enumerate(registros, start=2):
        # Consecutivo en columna A
        celda_a = ws.cell(row=fila_idx, column=1, value=fila_idx - 1)
        celda_a.font = FUENTE_DATO

        for col_idx, nombre_col in enumerate(columnas, start=2):
            valor = normalizar_valor(registro.get(nombre_col))
            celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
            celda.font = FUENTE_DATO

    ultima_fila = len(registros) + 1
    ultima_col = len(columnas) + 1
    ultima_col_letra = get_column_letter(ultima_col)

    # --- Anchos de columna ajustados al contenido ---
    autoajustar_anchos(ws, ultima_col, ultima_fila)

    # --- Paneles inmovilizados (igual que la plantilla: C2) ---
    ws.freeze_panes = "C2"

    # --- Autofiltro sobre todo el rango de datos ---
    ws.auto_filter.ref = f"A1:{ultima_col_letra}{ultima_fila}"

    wb.save(ruta_salida)
    print(f"Excel generado: {ruta_salida}")
    print(f"Registros escritos: {len(registros)}")
    print(f"Columnas: {len(columnas)}")


def main():
    ruta_json = "../data/perfiles.json"
    ruta_salida = f"../outputs/Perfiles.xlsx"

    if not os.path.exists(ruta_json):
        print(f"No se encontró el archivo: {ruta_json}")
        sys.exit(1)

    registros = cargar_registros(ruta_json)
    generar_excel(registros, ruta_salida)


if __name__ == "__main__":
    main()