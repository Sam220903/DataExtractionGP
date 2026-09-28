import json
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

print("Iniciando generación del reporte Excel...")

# ── Configuración ─────────────────────────────────────────────────────────────
YELLOW = "FFFF00"
WHITE = "FFFFFF"

def _fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

# ── Columnas del Excel ────────────────────────────────────────────────────────
HEADERS = ["Año Legislatura",  # A  col 1
    "Periodo",  # B  col 2
    "Folio periodo",  # C  col 3  — fondo amarillo
    "Folio legislatura",  # D  col 4  — fondo amarillo
    "Fecha",  # E  col 5
    "Tipo",  # F  col 6
    "Título",  # G  col 7
    "Tema 1",  # H  col 8  — fondo amarillo
    "Tema",  # I  col 9  — fondo amarillo
    "A favor",  # J  col 10
    "Contra",  # K  col 11
    "Abstenciones",  # L  col 12
    "Ausentes",  # M  col 13
    "Total",  # N  col 14  —
    "Porcentaje de Participación",  # O  col 15
    "Unidad",  # P  col 16
]

# Columnas con fondo amarillo en el encabezado (1-indexed)
YELLOW_HEADER_COLS = {3, 4, 8, 9}  # C, D, H, I

# Mapeo JSON key → índice de columna (1-indexed)
JSON_TO_COL = {"Año Legislatura": 1, "Periodo": 2, "Folio periodo": 3, "Folio legislatura": 4, "Fecha": 5, "Tipo": 6, "Título": 7, "Tema 1": 8, "Tema": 9, "A favor": 10,
    "Contra": 11, "Abstenciones": 12, "Ausentes": 13, }

# ── Carga de datos ────────────────────────────────────────────────────────────
print("  > Localizando y cargando datos JSON...")
CURRENT_FOLDER = os.path.dirname(os.path.abspath(__file__))
ruta_json = os.path.join(CURRENT_FOLDER, "..", "data", "unidad_participacion.json")

try:
    with open(ruta_json, encoding="utf-8") as f:
        data = json.load(f)
except FileNotFoundError:
    print(f"\n❌ ERROR: No se encontró el archivo JSON en la ruta calculada:")
    print(f"   {ruta_json}")
    print("   Por favor, verifica que la estructura de carpetas sea la correcta.")
    exit()

registros = data["registros"]
total_registros = len(registros)

# Ordenar los registros por fecha (del más antiguo al más reciente) y por después por folios
# Ordenamiento de múltiples niveles (Fecha -> Folio periodo -> Folio legislatura)
registros.sort(key=lambda x: (
    x.get("Fecha") or "",             # 1er nivel: Fecha
    x.get("Folio periodo") or 0,      # 2do nivel: Desempata por Folio periodo
    x.get("Folio legislatura") or 0   # 3er nivel: Desempata por Folio legislatura
))
print(f"  > ✓ Datos cargados correctamente ({total_registros} registros).")

# ── Creación del libro ────────────────────────────────────────────────────────
print("  > Estructurando documento Excel...")
wb = Workbook()
ws = wb.active
ws.title = "Hoja 1"

# ── Fila 1: encabezados ───────────────────────────────────────────────────────
for col_idx, header in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = Font(bold=True)
    if col_idx in YELLOW_HEADER_COLS:
        cell.fill = _fill(YELLOW)

# ── Filas de datos ────────────────────────────────────────────────────────────
print("  > Procesando filas y aplicando fórmulas...")
for row_offset, registro in enumerate(registros):
    row = 2 + row_offset
    
    # Retroalimentación no invasiva para archivos grandes (imprime cada 500 registros)
    if row_offset > 0 and row_offset % 500 == 0:
        print(f"    - Procesados {row_offset} de {total_registros}...")

    # Columnas A–M desde el JSON
    for key, col_idx in JSON_TO_COL.items():
        val = registro.get(key)

        # Convertir fecha de string a datetime si aplica
        if key == "Fecha" and isinstance(val, str):
            try:
                val = datetime.strptime(val, "%Y-%m-%d")
            except ValueError:
                pass

        cell = ws.cell(row=row, column=col_idx, value=val)

        if key == "Fecha":
            cell.number_format = "dd/mm/yyyy"

    # Columnas N, O, P: siempre fórmula para garantizar consistencia ───────────
    ws.cell(row=row, column=14, value=f"=SUM(J{row}:M{row})")

    o_cell = ws.cell(row=row, column=15, value=f"=SUM(J{row}:L{row})/SUM(J{row}:M{row})")
    o_cell.number_format = "0%"

    p_cell = ws.cell(row=row, column=16, value=f"=MAX(J{row}:L{row})/SUM(J{row}:L{row})")
    p_cell.number_format = "0.00%"

# ── Ancho de columnas ─────────────────────────────────────────────────────────
for col_idx in range(1, len(HEADERS) + 1):
    ws.column_dimensions[get_column_letter(col_idx)].width = 13.0

# ── Inmovilizar fila de encabezado ────────────────────────────────────────────
ws.freeze_panes = "A2"

# ── Guardar ───────────────────────────────────────────────────────────────────
print("  > Guardando el archivo final...")
output_filename = "Unidad y Participación.xlsx"

# Construcción de ruta absoluta para el guardado
output_folder = os.path.join(CURRENT_FOLDER, "..", "outputs")
os.makedirs(output_folder, exist_ok=True) # Crea la carpeta 'outputs' si no existe

output_path = os.path.join(output_folder, output_filename)
wb.save(output_path)

print(f"\n ¡Proceso finalizado con éxito!")
print(f"  Archivo generado en: {output_path}")